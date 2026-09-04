#!/usr/bin/env python3
"""Gera workbook de auditoria a partir do JSON emitido pelo gate."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


BLUE = PatternFill("solid", fgColor="1F4E79")
RED = PatternFill("solid", fgColor="F4CCCC")
GREEN = PatternFill("solid", fgColor="D9EAD3")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(color="FFFFFF", bold=True)


def _header(ws: Any, row: int, columns: list[str]) -> None:
    for col, value in enumerate(columns, 1):
        cell = ws.cell(row, col, value)
        cell.fill = BLUE
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fit(ws: Any, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def build_workbook(report: dict[str, Any], output: Path) -> None:
    wb = openpyxl.Workbook()
    resumo = wb.active
    resumo.title = "Resumo Executivo"
    resumo["A1"] = "GATE PRÉ-ENVIO — SQUAD FPE"
    resumo["A1"].font = Font(size=16, bold=True)
    resumo["A3"] = "Status"
    resumo["B3"] = report.get("gate", "BLOQUEADO")
    resumo["B3"].fill = GREEN if report.get("gate") == "PRE_LIBERADO" else RED
    resumo["A4"] = "Gerado em"
    resumo["B4"] = report.get("generated_at")
    resumo["A5"] = "Pasta auditada"
    resumo["B5"] = report.get("target")
    resumo["A7"] = "Achados"
    resumo["B7"] = len(report.get("findings", []))
    resumo.column_dimensions["A"].width = 24
    resumo.column_dimensions["B"].width = 90

    findings = wb.create_sheet("Ressalvas")
    findings["A1"] = "RESSALVAS E BLOQUEIOS"
    findings["A1"].font = Font(size=14, bold=True)
    _header(findings, 4, ["Item SUPAT", "Tipo", "Mensagem"])
    for row, finding in enumerate(report.get("findings", []), 5):
        findings.cell(row, 1, finding.get("item"))
        findings.cell(row, 2, finding.get("tipo"))
        findings.cell(row, 3, finding.get("mensagem"))
        for col in range(1, 4):
            findings.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            findings.cell(row, col).fill = RED
    _fit(findings, [18, 34, 100])

    item_sheet = wb.create_sheet("Auditoria Item a Item")
    item_sheet["A1"] = "VALIDAÇÃO MECÂNICA DE CÓDIGOS E PREÇOS"
    item_sheet["A1"].font = Font(size=14, bold=True)
    _header(item_sheet, 4, ["Linha", "Item", "Código", "Banco", "Unidade", "Custo direto", "Falhas"])
    items = report.get("code_audit", {}).get("itens", [])
    for row, item in enumerate(items, 5):
        values = [item.get("linha"), item.get("item"), item.get("codigo"), item.get("banco"), item.get("unidade"), item.get("custo_direto"), "; ".join(item.get("falhas", []))]
        for col, value in enumerate(values, 1):
            item_sheet.cell(row, col, value)
            item_sheet.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        item_sheet.cell(row, 7).fill = RED if item.get("falhas") else GREEN
    _fit(item_sheet, [10, 16, 18, 16, 12, 18, 80])

    checklist = wb.create_sheet("Simulação SUPAT ITEM 01-10")
    checklist["A1"] = "SIMULAÇÃO SUPAT / SAEB — ITEM 01 A 10"
    checklist["A1"].font = Font(size=14, bold=True)
    _header(checklist, 4, ["Item", "Status", "Achados"])
    by_item: dict[str, list[str]] = {f"ITEM {number:02d}": [] for number in range(1, 11)}
    for finding in report.get("findings", []):
        item = finding.get("item", "GERAL")
        by_item.setdefault(item, []).append(f"{finding.get('tipo')}: {finding.get('mensagem')}")
    for row, (item, messages) in enumerate(by_item.items(), 5):
        checklist.cell(row, 1, item)
        checklist.cell(row, 2, "BLOQUEADO" if messages else "CONFORME")
        checklist.cell(row, 3, "\n".join(messages) if messages else "Sem achado registrado")
        checklist.cell(row, 2).fill = RED if messages else GREEN
        checklist.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    _fit(checklist, [18, 22, 110])

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("audit_json", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    with args.audit_json.open("r", encoding="utf-8") as stream:
        report = json.load(stream)
    build_workbook(report, args.output)
    print(f"[OK] Workbook gerado: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
