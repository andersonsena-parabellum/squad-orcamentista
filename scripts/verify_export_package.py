#!/usr/bin/env python3
"""Verifica o pacote exportado que sera efetivamente selado."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

from amostrar_codigos import PROCESS_NOTE_RE
from squad_common import load_json, normalize_text, sha256_file, utc_now


def _pdf_text(path: Path) -> str:
    try:
        import pypdf
    except ImportError as exc:
        raise RuntimeError("Dependencia pypdf ausente; verificacao PDF e obrigatoria") from exc
    reader = pypdf.PdfReader(path)
    if not reader.pages:
        raise RuntimeError("PDF sem paginas")
    return normalize_text(" ".join(page.extract_text() or "" for page in reader.pages))


def _xlsx_descriptions(path: Path) -> tuple[list[str], list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    failures = []
    descriptions = []
    for ws in wb.worksheets:
        description_col = None
        header_row = None
        for row in range(1, min(ws.max_row, 30) + 1):
            for col in range(1, min(ws.max_column, 30) + 1):
                if "DESCRI" in normalize_text(ws.cell(row, col).value):
                    description_col, header_row = col, row
                    break
            if description_col:
                break
        if not description_col:
            continue
        for row in range(header_row + 1, ws.max_row + 1):
            description = str(ws.cell(row, description_col).value or "").strip()
            if not description:
                continue
            descriptions.append(description)
            if PROCESS_NOTE_RE.search(description):
                failures.append(f"Descricao contaminada em {ws.title}!{ws.cell(row, description_col).coordinate}")
    return descriptions, failures


def verify(package_dir: Path, state_path: Path) -> dict[str, Any]:
    findings = []
    artifacts = []
    state = load_json(state_path)
    if state.get("status_fluxo") != "EXPORTACAO_STAGING":
        findings.append({"type": "ESTADO_INVALIDO", "detail": f"Esperado EXPORTACAO_STAGING; atual={state.get('status_fluxo')}"})
    if state.get("status_gate") != "PRE_LIBERADO":
        findings.append({"type": "GATE_TECNICO_INVALIDO", "detail": state.get("status_gate")})

    for number in range(1, 8):
        dirs = [path for path in package_dir.iterdir() if path.is_dir() and re.match(rf"^0{number}(?:-|$)", path.name)]
        if len(dirs) != 1:
            findings.append({"type": "PASTA_OFICIAL_INVALIDA", "detail": f"Prefixo 0{number}: {len(dirs)} pastas"})
            continue
        folder = dirs[0]
        xlsx_files = [path for path in folder.glob("*.xlsx") if not path.name.startswith("~$")]
        pdf_files = list(folder.glob("*.pdf"))
        if len(xlsx_files) != 1 or len(pdf_files) != 1:
            findings.append({"type": "PAR_XLSX_PDF_INVALIDO", "detail": f"{folder.name}: xlsx={len(xlsx_files)}, pdf={len(pdf_files)}"})
            continue
        xlsx, pdf = xlsx_files[0], pdf_files[0]
        if xlsx.stat().st_size == 0 or pdf.stat().st_size == 0:
            findings.append({"type": "ARQUIVO_VAZIO", "detail": folder.name})
            continue
        descriptions, xlsx_failures = _xlsx_descriptions(xlsx)
        findings.extend({"type": "DESCRICAO_CONTAMINADA", "detail": message} for message in xlsx_failures)
        try:
            pdf_text = _pdf_text(pdf)
        except Exception as exc:
            findings.append({"type": "PDF_INVALIDO", "detail": f"{pdf}: {exc}"})
            pdf_text = ""
        for description in descriptions[:5]:
            probe = normalize_text(description)[:30]
            if probe and probe not in pdf_text:
                findings.append({"type": "XLSX_PDF_DIVERGENTE", "detail": f"Trecho nao localizado no PDF: {probe}"})
        artifacts.extend(
            [
                {"path": str(xlsx.resolve()), "sha256": sha256_file(xlsx), "bytes": xlsx.stat().st_size},
                {"path": str(pdf.resolve()), "sha256": sha256_file(pdf), "bytes": pdf.stat().st_size},
            ]
        )
    return {"schema_version": "1.0.0", "generated_at": utc_now(), "package": str(package_dir.resolve()), "gate": "POS_EXPORTACAO_OK" if not findings else "BLOQUEADO", "findings": findings, "artifacts": artifacts}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pacote", type=Path)
    parser.add_argument("--estado", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if not args.pacote.is_dir() or not args.estado.is_file():
        print("[BLOQUEADO] Pacote ou estado inexistente", file=sys.stderr)
        return 2
    report = verify(args.pacote.resolve(), args.estado.resolve())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"gate": report["gate"], "findings": len(report["findings"]), "out": str(args.out)}, ensure_ascii=False))
    return 0 if report["gate"] == "POS_EXPORTACAO_OK" else 2


if __name__ == "__main__":
    raise SystemExit(main())
