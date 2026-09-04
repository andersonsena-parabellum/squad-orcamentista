#!/usr/bin/env python3
"""Importa uma publicação mensal oficial do SINAPI para a base local.

O importador trabalha diretamente com o ZIP XLSX publicado pela CAIXA,
confere hash, competência, UF e consistência entre os relatórios. A troca dos
dados SINAPI no SQLite é atômica e não altera fontes de outros bancos.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sqlite3
import tempfile
import unicodedata
import zipfile
from contextlib import closing
from pathlib import Path
from typing import Any

import openpyxl

from squad_common import PRICE_DB, SOURCE_REGISTRY, atomic_write_json, exclusive_lock, load_json, sha256_file, utc_now


DEFAULT_SOURCE = SOURCE_REGISTRY.parent / "00-FONTES_OFICIAIS" / "SINAPI-2026-07-formato-xlsx.zip"
DEFAULT_URL = "https://www.caixa.gov.br/Downloads/sinapi-relatorios-mensais/SINAPI-2026-07-formato-xlsx.zip"
DEFAULT_SHA256 = "58c131f997560332cf2d7f7f90644790d5c6e2a909a2963b18f136779312b14f"
CODE_IN_FORMULA_RE = re.compile(r",\s*(\d+)\s*\)\s*$")


def ascii_fold(value: str) -> str:
    return "".join(char for char in unicodedata.normalize("NFKD", value) if not unicodedata.combining(char))


def normalized_code(value: Any) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return str(int(value))
    text = str(value).strip()
    if text.startswith("="):
        match = CODE_IN_FORMULA_RE.search(text)
        return match.group(1) if match else None
    return text if text.isdigit() else None


def positive_or_none(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
        return float(value)
    return None


def find_reference_workbook(archive: zipfile.ZipFile) -> str:
    candidates = [
        name
        for name in archive.namelist()
        if name.lower().endswith(".xlsx") and "_referencia_" in ascii_fold(name).lower()
    ]
    if len(candidates) != 1:
        raise RuntimeError(f"ZIP deve conter exatamente um SINAPI_Referência_*.xlsx; encontrados: {candidates}")
    return candidates[0]


def find_uf_column(sheet: Any, uf: str, uf_row: int, header_row: int, expected_header: str) -> int:
    for column in range(1, sheet.max_column + 1):
        uf_value = ascii_fold(str(sheet.cell(uf_row, column).value or "")).strip().upper()
        header = ascii_fold(str(sheet.cell(header_row, column).value or "")).strip().upper()
        if uf_value == uf and expected_header in header:
            return column
    raise RuntimeError(f"Coluna {uf}/{expected_header} não encontrada na aba {sheet.title}")


def read_price_sheet(sheet: Any, uf: str, kind: str) -> tuple[str, dict[str, dict[str, Any]]]:
    if kind == "composition":
        uf_column = find_uf_column(sheet, uf, 9, 10, "CUSTO")
    else:
        uf_column = find_uf_column(sheet, uf, 10, 10, uf)
    competence = str(sheet.cell(3, 2).value or "").strip()
    records: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=11, values_only=True):
        code = normalized_code(row[1])
        if not code:
            continue
        if code in records:
            raise RuntimeError(f"Código duplicado em {sheet.title}: {code}")
        records[code] = {
            "grupo_ou_tipo": str(row[0] or "").strip(),
            "descricao": str(row[2] or "").strip(),
            "unidade": str(row[3] or "").strip(),
            "preco": positive_or_none(row[uf_column - 1]),
        }
    return competence, records


def reconcile_regimes(
    without_relief: dict[str, dict[str, Any]],
    with_relief: dict[str, dict[str, Any]],
    label: str,
) -> list[dict[str, Any]]:
    if set(without_relief) != set(with_relief):
        missing_a = sorted(set(with_relief) - set(without_relief))[:10]
        missing_b = sorted(set(without_relief) - set(with_relief))[:10]
        raise RuntimeError(f"Códigos divergentes em {label}: só desonerado={missing_a}; só não desonerado={missing_b}")
    output: list[dict[str, Any]] = []
    for code in sorted(without_relief, key=lambda item: int(item)):
        nao_deson = without_relief[code]
        deson = with_relief[code]
        for field in ("descricao", "unidade", "grupo_ou_tipo"):
            if nao_deson[field] != deson[field]:
                raise RuntimeError(f"{label} {code}: campo {field} diverge entre regimes")
        output.append(
            {
                "codigo": code,
                **{field: nao_deson[field] for field in ("descricao", "unidade", "grupo_ou_tipo")},
                "nao_deson": nao_deson["preco"],
                "deson": deson["preco"],
            }
        )
    return output


def read_analytical(sheet: Any) -> tuple[dict[str, dict[str, Any]], list[tuple[Any, ...]], dict[str, dict[str, Any]]]:
    headers: dict[str, dict[str, Any]] = {}
    items: list[tuple[Any, ...]] = []
    missing_inputs: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=11, values_only=True):
        parent = normalized_code(row[1])
        if not parent:
            continue
        item_type = str(row[2] or "").strip().upper()
        item_code = normalized_code(row[3])
        description = str(row[4] or "").strip()
        unit = str(row[5] or "").strip()
        coefficient = row[6]
        situation = str(row[7] or "").strip().upper()
        if not item_type and not item_code:
            if parent in headers:
                raise RuntimeError(f"Cabeçalho analítico duplicado: {parent}")
            headers[parent] = {
                "grupo": str(row[0] or "").strip(),
                "descricao": description,
                "unidade": unit,
                "situacao": situation,
            }
            continue
        if item_type not in {"INSUMO", "COMPOSICAO"} or not item_code:
            raise RuntimeError(f"Item analítico inválido na composição {parent}: tipo={item_type!r}, código={item_code!r}")
        if not isinstance(coefficient, (int, float)) or isinstance(coefficient, bool) or coefficient < 0:
            raise RuntimeError(f"Coeficiente inválido: composição {parent}, item {item_code}, valor={coefficient!r}")
        items.append(("SINAPI", parent, item_type, item_code, description, unit, float(coefficient)))
        if item_type == "INSUMO":
            missing_inputs.setdefault(item_code, {"descricao": description, "unidade": unit, "situacao": situation})
    return headers, items, missing_inputs


def parse_publication(source: Path, uf: str) -> dict[str, Any]:
    with zipfile.ZipFile(source) as archive, tempfile.TemporaryDirectory(prefix="sinapi-import-") as temp_dir:
        workbook_name = find_reference_workbook(archive)
        archive.extract(workbook_name, temp_dir)
        workbook_path = Path(temp_dir) / workbook_name
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            comp_nao_meta, comp_nao = read_price_sheet(workbook["CSD"], uf, "composition")
            comp_des_meta, comp_des = read_price_sheet(workbook["CCD"], uf, "composition")
            input_nao_meta, input_nao = read_price_sheet(workbook["ISD"], uf, "input")
            input_des_meta, input_des = read_price_sheet(workbook["ICD"], uf, "input")
            competencies = {comp_nao_meta, comp_des_meta, input_nao_meta, input_des_meta}
            if len(competencies) != 1:
                raise RuntimeError(f"Competências divergentes no arquivo: {sorted(competencies)}")
            composition_rows = reconcile_regimes(comp_nao, comp_des, "composição")
            input_rows = reconcile_regimes(input_nao, input_des, "insumo")
            analytical_headers, analytical_items, analytical_inputs = read_analytical(workbook["Analítico"])
        finally:
            workbook.close()

    composition_codes = {row["codigo"] for row in composition_rows}
    if composition_codes != set(analytical_headers):
        raise RuntimeError("O conjunto de composições dos relatórios de custo diverge do relatório analítico")
    for row in composition_rows:
        header = analytical_headers[row["codigo"]]
        if row["descricao"] != header["descricao"] or row["unidade"] != header["unidade"]:
            raise RuntimeError(f"Descrição/unidade diverge no analítico para composição {row['codigo']}")
    priced_input_codes = {row["codigo"] for row in input_rows}
    extra_input_rows = [
        {
            "codigo": code,
            "descricao": record["descricao"],
            "unidade": record["unidade"],
            "grupo_ou_tipo": "SEM PREÇO",
            "nao_deson": None,
            "deson": None,
        }
        for code, record in analytical_inputs.items()
        if code not in priced_input_codes
    ]
    month, year = next(iter(competencies)).split("/")
    return {
        "competencia": f"{year}-{month}",
        "composicoes": composition_rows,
        "insumos": input_rows + sorted(extra_input_rows, key=lambda row: int(row["codigo"])),
        "itens": analytical_items,
        "sem_preco": {
            "composicoes": sum(1 for row in composition_rows if not row["deson"] or not row["nao_deson"]),
            "insumos": sum(1 for row in input_rows + extra_input_rows if not row["deson"] or not row["nao_deson"]),
        },
    }


def import_database(data: dict[str, Any]) -> None:
    if not PRICE_DB.is_file():
        raise FileNotFoundError(f"Banco ausente: {PRICE_DB}")
    fd, temp_name = tempfile.mkstemp(prefix=".base_precos.", suffix=".db", dir=PRICE_DB.parent)
    os.close(fd)
    temp_db = Path(temp_name)
    try:
        shutil.copy2(PRICE_DB, temp_db)
        with closing(sqlite3.connect(temp_db)) as connection:
            connection.execute("PRAGMA foreign_keys=ON")
            connection.execute("BEGIN IMMEDIATE")
            connection.execute("DELETE FROM composicao_itens WHERE fonte='SINAPI'")
            connection.execute("DELETE FROM composicoes WHERE fonte='SINAPI'")
            connection.execute("DELETE FROM insumos WHERE fonte='SINAPI'")
            connection.executemany(
                "INSERT INTO composicoes(fonte,codigo,descricao,unidade,grupo,custo_deson,custo_nao_deson) "
                "VALUES('SINAPI',?,?,?,?,?,?)",
                [
                    (row["codigo"], row["descricao"], row["unidade"], row["grupo_ou_tipo"], row["deson"], row["nao_deson"])
                    for row in data["composicoes"]
                ],
            )
            connection.executemany(
                "INSERT INTO insumos(fonte,codigo,descricao,unidade,tipo,preco_deson,preco_nao_deson) "
                "VALUES('SINAPI',?,?,?,?,?,?)",
                [
                    (row["codigo"], row["descricao"], row["unidade"], row["grupo_ou_tipo"], row["deson"], row["nao_deson"])
                    for row in data["insumos"]
                ],
            )
            connection.executemany(
                "INSERT INTO composicao_itens(fonte,codigo_composicao,tipo_item,codigo_item,descricao,unidade,coeficiente) "
                "VALUES(?,?,?,?,?,?,?)",
                data["itens"],
            )
            metadata = {
                "sinapi_competencia": data["competencia"],
                "sinapi_uf": data["uf"],
                "sinapi_sha256_origem": data["sha256_origem"],
                "sinapi_imported_at": utc_now(),
            }
            connection.executemany("INSERT OR REPLACE INTO metadata(chave,valor) VALUES(?,?)", metadata.items())
            connection.commit()
            integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
            if integrity != "ok":
                raise RuntimeError(f"SQLite não passou no integrity_check: {integrity}")
            orphans = connection.execute(
                "SELECT count(*) FROM composicao_itens ci LEFT JOIN insumos i "
                "ON i.fonte=ci.fonte AND i.codigo=ci.codigo_item "
                "WHERE ci.fonte='SINAPI' AND ci.tipo_item='INSUMO' AND i.codigo IS NULL"
            ).fetchone()[0]
            orphans += connection.execute(
                "SELECT count(*) FROM composicao_itens ci LEFT JOIN composicoes c "
                "ON c.fonte=ci.fonte AND c.codigo=ci.codigo_item "
                "WHERE ci.fonte='SINAPI' AND ci.tipo_item='COMPOSICAO' AND c.codigo IS NULL"
            ).fetchone()[0]
            if orphans:
                raise RuntimeError(f"Importação produziria {orphans} referências órfãs")
        os.replace(temp_db, PRICE_DB)
    finally:
        if temp_db.exists():
            temp_db.unlink()


def update_registry(source: Path, source_hash: str, url: str, data: dict[str, Any]) -> None:
    registry = load_json(SOURCE_REGISTRY)
    try:
        relative_source = source.resolve().relative_to(SOURCE_REGISTRY.parent.resolve()).as_posix()
    except ValueError as exc:
        raise RuntimeError("O arquivo oficial deve estar dentro de 03-BASE_DE_PRECOS") from exc
    registry["database"]["sha256"] = sha256_file(PRICE_DB)
    registry["fontes"]["SINAPI"] = {
        "status": "LIBERADA",
        "uf": data["uf"],
        "competencia": data["competencia"],
        "regimes": ["DESONERADO", "NAO_DESONERADO"],
        "arquivo_origem": relative_source,
        "url_oficial": url,
        "sha256_origem": source_hash,
        "importado_em": utc_now(),
        "contagens": {
            "composicoes": len(data["composicoes"]),
            "insumos": len(data["insumos"]),
            "itens_analiticos": len(data["itens"]),
            "sem_preco_na_uf": data["sem_preco"],
        },
        "motivo": "Publicação mensal oficial da CAIXA conferida por hash e importada de forma transacional. Registros sem preço permanecem bloqueados individualmente.",
    }
    atomic_write_json(SOURCE_REGISTRY, registry)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--arquivo", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--uf", default="BA")
    parser.add_argument("--url", default=DEFAULT_URL)
    parser.add_argument("--sha256-esperado", default=DEFAULT_SHA256)
    args = parser.parse_args()
    source = args.arquivo.resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Arquivo oficial ausente: {source}")
    actual_hash = sha256_file(source)
    if actual_hash.lower() != args.sha256_esperado.lower():
        raise RuntimeError(f"Hash divergente: esperado={args.sha256_esperado}; obtido={actual_hash}")
    data = parse_publication(source, args.uf.upper())
    data["uf"] = args.uf.upper()
    data["sha256_origem"] = actual_hash
    lock = PRICE_DB.with_suffix(".import.lock")
    with exclusive_lock(lock, "importar_sinapi"):
        import_database(data)
        update_registry(source, actual_hash, args.url, data)
    print(
        json.dumps(
            {
                "status": "LIBERADA",
                "fonte": "SINAPI",
                "uf": data["uf"],
                "competencia": data["competencia"],
                "composicoes": len(data["composicoes"]),
                "insumos": len(data["insumos"]),
                "itens_analiticos": len(data["itens"]),
                "sem_preco_na_uf": data["sem_preco"],
                "sha256_origem": actual_hash,
                "sha256_database": sha256_file(PRICE_DB),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
