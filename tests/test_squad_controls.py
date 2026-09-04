from __future__ import annotations

import argparse
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from atualizar_estado import update
from amostrar_codigos import audit_workbook
from simulate_supat_preenvio import REQUIRED_HUMAN_ITEMS, audit_quotes, human_checkpoints, valid_cnpj
from squad_common import atomic_write_json, sha256_file, source_is_releasable, validate_handoff_file
from validate_repository import validate as validate_repository


def make_cnpj(base12: str) -> str:
    numbers = [int(char) for char in base12]
    for weights in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        remainder = sum(number * weight for number, weight in zip(numbers, weights)) % 11
        numbers.append(0 if remainder < 2 else 11 - remainder)
    return "".join(str(number) for number in numbers)


class SquadControlTests(unittest.TestCase):
    def test_quarantined_source_is_fail_closed(self) -> None:
        released, reason = source_is_releasable("ORSE")
        self.assertFalse(released)
        self.assertIn("QUARENTENADA", reason)

    def test_official_sinapi_source_is_releasable(self) -> None:
        released, reason = source_is_releasable("SINAPI")
        self.assertTrue(released, reason)

    def test_cnpj_checksum(self) -> None:
        self.assertTrue(valid_cnpj("11.222.333/0001-81"))
        self.assertFalse(valid_cnpj("11.111.111/1111-11"))

    def test_quote_map_rejects_arithmetic_mean_as_median(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "DC-007.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Item", "Especificação", "Fornecedor", "CNPJ", "Data", "Preço FOB", "Mediana"])
            for index, price in enumerate((100, 200, 900), 1):
                ws.append(["1", "Rack 44U 800x1000", f"Fornecedor {index}", make_cnpj(f"12345678{index:04d}"), date.today(), price, 400])
            wb.save(path)
            findings = audit_quotes(path)
            self.assertTrue(any(item["tipo"] == "MEDIANA_DIVERGENTE" for item in findings))

    def test_formula_without_cached_value_blocks_budget(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "DC-001.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Item", "Código", "Banco", "Descrição", "Und", "Quant.", "Valor Unit", "Valor Unit com BDI", "Total", "Peso %"])
            ws.append(["1", "93680", "SINAPI", "EXECUÇÃO DE PAVIMENTO EM PISO INTERTRAVADO, COM BLOCO RETANGULAR COLORIDO DE 20 X 10 CM, ESPESSURA 6 CM. AF_10/2022", "M2", 1, 80.04, 100, "=F2*H2", 100])
            wb.save(path)
            report = audit_workbook(path, "NAO_DESONERADO")
            types = {failure["tipo"] for failure in report["falhas"]}
            self.assertTrue(any(item.startswith("FORMULA_SEM_CACHE") for item in types))
            self.assertEqual(report["gate"], "BLOQUEADO")

    def test_official_code_without_ba_price_stays_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "DC-001.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Item", "Código", "Banco", "Descrição", "Und", "Quant.", "Valor Unit", "Valor Unit com BDI", "Total", "Peso %"])
            ws.append([
                "1",
                "105006",
                "SINAPI",
                "RAMPA DE ACESSIBILIDADE EM CONCRETO PRÉ MOLDADO, EM CALÇADA NOVA COM LARGURA MAIOR OU IGUAL À 3,00 M, FCK 25MPA, COM PISO PODOTÁTIL. AF_03/2024",
                "UN",
                1,
                1,
                1,
                1,
                100,
            ])
            wb.save(path)
            report = audit_workbook(path, "NAO_DESONERADO")
            types = {failure["tipo"] for failure in report["falhas"]}
            self.assertIn("PRECO_OFICIAL_ZERO_OU_AUSENTE", types)
            self.assertEqual(report["gate"], "BLOQUEADO")

    def test_repository_is_operational_with_quarantines_isolated(self) -> None:
        report = validate_repository()
        self.assertEqual(report["gate"], "LIBERADO")
        self.assertGreater(report["restricoes_ativas"], 0)

    def test_handoff_rejects_hash_drift(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            source = project / "memoria.xlsx"
            source.write_bytes(b"original")
            state = {
                "project_id": "p1",
                "revision_id": "R00",
                "manifesto_fontes": {"sha256": "a" * 64},
            }
            handoff = project / "handoff.json"
            atomic_write_json(
                handoff,
                {
                    "schema_version": "1.0.0",
                    "project_id": "p1",
                    "revision_id": "R00",
                    "producer": "levi",
                    "consumer": "elias",
                    "created_at": "2026-09-04T00:00:00+00:00",
                    "source_manifest_sha256": "a" * 64,
                    "artifacts": [{"path": source.name, "sha256": sha256_file(source)}],
                    "status": "PRONTO",
                    "open_issues": [],
                },
            )
            source.write_bytes(b"alterado")
            with self.assertRaisesRegex(ValueError, "Hash divergente"):
                validate_handoff_file(handoff, state, "levi", project)

    def test_wrong_actor_cannot_advance_state(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            squad = project / "_squad"
            squad.mkdir()
            manifest = squad / "manifesto_fontes.json"
            atomic_write_json(manifest, {"files": []})
            atomic_write_json(
                squad / "estado.json",
                {
                    "status_fluxo": "FONTES_CONGELADAS",
                    "manifesto_fontes": {"path": manifest.name, "sha256": sha256_file(manifest)},
                    "artifacts": {},
                },
            )
            args = argparse.Namespace(
                pasta_obra=project,
                agente="otavio",
                transicao="QTO_PRONTO",
                artefato=None,
                justificativa=None,
            )
            with self.assertRaises(PermissionError):
                update(args)

    def test_missing_human_checkpoints_never_passes_by_omission(self) -> None:
        findings = human_checkpoints(None)
        self.assertEqual(len(findings), len(REQUIRED_HUMAN_ITEMS))
        self.assertEqual(REQUIRED_HUMAN_ITEMS, {f"ITEM {number:02d}" for number in range(1, 11)})

    def test_valid_handoff_advances_only_authorized_phase(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            squad = project / "_squad"
            squad.mkdir()
            source = project / "projeto.pdf"
            source.write_bytes(b"fonte")
            deliverable = project / "memoria.xlsx"
            deliverable.write_bytes(b"qto")
            manifest = squad / "manifesto_fontes.json"
            atomic_write_json(manifest, {"files": [{"relative_path": source.name, "sha256": sha256_file(source)}]})
            state = {
                "project_id": "p1",
                "revision_id": "R00",
                "status_fluxo": "FONTES_CONGELADAS",
                "status_gate": "EM_ANDAMENTO",
                "manifesto_fontes": {"path": manifest.name, "sha256": sha256_file(manifest)},
                "artifacts": {},
            }
            atomic_write_json(squad / "estado.json", state)
            handoff = squad / "handoff-levi.json"
            atomic_write_json(
                handoff,
                {
                    "schema_version": "1.0.0",
                    "project_id": "p1",
                    "revision_id": "R00",
                    "producer": "levi",
                    "consumer": "elias",
                    "created_at": "2026-09-04T00:00:00+00:00",
                    "source_manifest_sha256": sha256_file(manifest),
                    "artifacts": [{"path": deliverable.name, "sha256": sha256_file(deliverable)}],
                    "status": "PRONTO",
                    "open_issues": [],
                },
            )
            args = argparse.Namespace(pasta_obra=project, agente="levi", transicao="QTO_PRONTO", artefato=handoff, justificativa=None)
            update(args)
            updated = json.loads((squad / "estado.json").read_text(encoding="utf-8"))
            self.assertEqual(updated["status_fluxo"], "QTO_PRONTO")
            self.assertIn("QTO_PRONTO", updated["artifacts"])


if __name__ == "__main__":
    unittest.main()
